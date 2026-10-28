"""The BOM as an Excel workbook (auto3d/bomxlsx.py, POST /api/bom/xlsx): a sheet per view, the header bold and frozen,
masses with two decimals, a structured BOM's names indented by their level; without openpyxl the server says so."""
import io
import unittest
from unittest import mock

from auto3d import bomxlsx

SUMMARY = {"title": "Souhrnný", "keys": ["pos", "name", "qty", "material", "mass_g", "how"],
           "head": ["Poz.", "Název", "Ks", "Materiál", "Hmotnost/ks (g)", "Výroba / nákup"],
           "rows": [[1, "Base", 1, "aluminium", 102.234, "vyráběný (obrábění)"], [3, "Pin", 7, "PLA", 0.41, "vyráběný (3D tisk)"],
                    [5, "M3×10 screw", 6, "steel", None, "nakupovaný"]]}
STRUCTURED = {"title": "Strukturovaný", "keys": ["pos", "level", "name", "qty", "total"], "head": ["Poz.", "Úroveň", "Název", "Ks v nadřazené", "Ks celkem"],
              "rows": [[1, 1, "Base", 1, 1], [2, 1, "Lid unit", 2, 2], [4, 2, "Lid", 1, 2], [3, 2, "Pin", 3, 6]], "levels": [0, 0, 1, 1]}


@unittest.skipUnless(bomxlsx.available(), "openpyxl not installed")
class TestBomXlsx(unittest.TestCase):
    def test_workbook(self):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(bomxlsx.to_xlsx([SUMMARY, STRUCTURED], "Machine — BOM")))
        self.assertEqual(wb.sheetnames, ["Souhrnný", "Strukturovaný"])
        ws = wb["Souhrnný"]
        self.assertEqual([c.value for c in ws[1]], SUMMARY["head"])
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual([c.value for c in ws[3]], [3, "Pin", 7, "PLA", 0.41, "vyráběný (3D tisk)"])
        self.assertEqual(ws["E2"].number_format, "0.00")
        self.assertIsNone(ws["E4"].value)  # (a mass not known: empty)
        st = wb["Strukturovaný"]
        self.assertEqual(st["C4"].alignment.indent, 2)  # Lid: inside Lid unit
        self.assertEqual(st["C2"].alignment.indent, 0)
        self.assertTrue(st["C3"].font.bold)  # Lid unit: its parts follow
        self.assertEqual(sum(st.cell(row=r, column=5).value for r in range(2, 6) if st.cell(row=r, column=3).value != "Lid unit"), 9)

    def test_endpoint(self):
        import server
        c = server.app.test_client()
        self.assertEqual(c.get("/api/bom/xlsx").get_json(), {"available": True})
        r = c.post("/api/bom/xlsx", json={"name": "Machine — BOM", "sheets": [SUMMARY]})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"PK"))  # (a zip: the .xlsx)
        self.assertIn("Machine - BOM.xlsx", r.headers["Content-Disposition"])


class TestWithoutOpenpyxl(unittest.TestCase):
    def test_csv_only(self):
        import server
        with mock.patch.object(bomxlsx, "_OK", False):
            c = server.app.test_client()
            self.assertEqual(c.get("/api/bom/xlsx").get_json(), {"available": False})
            r = c.post("/api/bom/xlsx", json={"name": "x", "sheets": [SUMMARY]})
            self.assertEqual(r.status_code, 501)
            self.assertIn("CSV", r.get_json()["error"])


if __name__ == "__main__":
    unittest.main()
