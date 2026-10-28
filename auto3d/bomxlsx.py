"""A bill of materials as an Excel workbook (.xlsx): a sheet per view (Summary, Structured), the header bold over a
rule and frozen, a filter on it, the columns as wide as their texts, masses with two decimals, a structured BOM's names
indented by their level. openpyxl is optional (requirements-modeler.txt): without it the app offers CSV only."""
import importlib
import io

_OK = None


def available() -> bool:
    """True when openpyxl can be imported."""
    global _OK
    if _OK is None:
        try:
            importlib.import_module("openpyxl")
            _OK = True
        except Exception:  # noqa: BLE001 — missing, or a broken install: either way not usable
            _OK = False
    return _OK


def to_xlsx(sheets, title="BOM") -> bytes:
    """sheets: [{title, head: [texts], rows: [[cells]], levels: [indent per row] (optional), keys: [column keys]
    (optional: "name" is indented, "mass_g" / "mass_total_g" get two decimals)}] -> the workbook's bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = str(title)
    rule = Border(bottom=Side(style="thin", color="888888"))
    for sh in sheets or []:
        ws = wb.create_sheet(str(sh.get("title") or "BOM")[:31])
        head, rows = list(sh.get("head") or []), list(sh.get("rows") or [])
        keys, levels = list(sh.get("keys") or []), list(sh.get("levels") or [])
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="EEF1F5")
            c.border = rule
        for r in rows:
            ws.append(["" if v is None else v for v in r])
        name_col = keys.index("name") + 1 if "name" in keys else None
        for i in range(len(rows)):
            if name_col and i < len(levels) and levels[i]:
                ws.cell(row=i + 2, column=name_col).alignment = Alignment(indent=2 * int(levels[i]))
            if name_col and i + 1 < len(levels) and levels[i + 1] > levels[i]:  # (an assembly: its parts follow)
                ws.cell(row=i + 2, column=name_col).font = Font(bold=True)
        for k, key in enumerate(keys, start=1):
            if key in ("mass_g", "mass_total_g"):
                for row in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                    row[0].number_format = "0.00"
        for k in range(1, len(head) + 1):
            width = max([len(str(head[k - 1]))] + [len(str(r[k - 1])) + (2 * levels[i] if k == name_col and i < len(levels) else 0)
                                                    for i, r in enumerate(rows) if k - 1 < len(r)])
            ws.column_dimensions[get_column_letter(k)].width = min(60, width + 2)
        ws.freeze_panes = "A2"
        if head:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"
    if not wb.sheetnames:
        wb.create_sheet("BOM")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
