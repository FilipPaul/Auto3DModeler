"""📋 The bill of materials, its whole hierarchy rolled up, in headless Edge. Parts Base, Lid, Pin, Bracket and a bought
ISO 4762 M3×10 screw; "Lid unit" (Lid, Pin ×3, screw ×2), "Hinge" (Bracket, Pin ×1); "Machine" (Base, Lid unit ×2,
Hinge, screw ×4 placed directly in it). -> get_bom summary: Base 1, Lid 2, Pin 7 (3 × 2 + 1), Bracket 1, screw 8 (4 + 2
× 2), make / buy, `from`; structured: the tree with qty per parent and the running total; Lid unit alone: Lid 1, Pin 3,
screw 2; get_assembly parts_list `total` -> 📋 BOM in the panel: Summary / Structured, a sub-assembly folded, another
assembly of the hierarchy chosen -> ⬇ CSV (its rows) and ⬇ Excel (.xlsx) -> in Czech: the headers -> the Machine's
sheet: 📋 Add BOM page -> 📚 All pages shows it -> the drawing set's PDF has the BOM page, its positions the balloons'
numbers -> a long BOM flows onto two pages. In Czech: the window's headers, materials (hliník, not aluminium) and make / buy,
the CSV's, the toasts (BOM saved, a library part put in), the BOM page and the assembly pages' parts lists.

  PYTHONIOENCODING=utf-8 py tools/e2e_bom.py http://127.0.0.1:5199 <out dir>
"""
import io
import json
import os
import re
import sys
import time

import fitz
from selenium import webdriver
from selenium.webdriver.common.by import By

URL, OUT = sys.argv[1], sys.argv[2]
opts = webdriver.EdgeOptions()
opts.add_argument("--headless=new")
opts.add_argument("--window-size=1600,1000")
opts.add_argument("--lang=en-US")
opts.set_capability("ms:loggingPrefs", {"browser": "ALL"})
d = webdriver.Edge(options=opts)
js = d.execute_script
A = "window.__app.assembly"
H = "window.__app.agentHandlers"
bad = []


def wait(script, t=90):
    end = time.time() + t
    while time.time() < end:
        v = js(script)
        if v:
            return v
        time.sleep(0.1)
    raise TimeoutError(script)


def call(t, **a):
    return js(f"return {H}[arguments[0]](arguments[1]).then(r => r.text).catch(e => 'ERROR ' + e.message)", t, a)


def check(ok, what):
    print(("  ok   " if ok else "  FAIL ") + what)
    if not ok:
        bad.append(what)


def block(name, w, h, t):
    call("new_model", name=name)
    call("add_feature", type="sketch", plane="XY", shapes=[{"rect": {"x": 0, "y": 0, "w": w, "h": h}}])
    call("add_feature", type="extrude", sketch="F1", d=t)


def table():
    """The 📋 BOM window's rows as texts per cell."""
    return js("return [...document.querySelectorAll('.bom-win tbody tr')].map(r => [...r.cells].map(c => c.innerText.trim()))")


SCREW = "ISO 4762 M3×10"
EN_MATS = {"aluminium", "steel", "brass", "stainless steel", "made", "assembly"}  # (English texts that must not show in Czech)
try:
    d.get(URL)
    wait("return !!window.__app")
    js("localStorage.removeItem('auto3d-lang'); localStorage.removeItem('a3d.bomMode')")
    # every toast's text as it appears (they go after a few seconds)
    js("window.__toasts = []; new MutationObserver((ms) => ms.forEach((m) => m.addedNodes.forEach((n) => window.__toasts.push(n.textContent)))).observe(document.getElementById('toasts'), {childList: true})")
    js("""window.__saved = []; window.showSaveFilePicker = async (o) => ({ name: o.suggestedName,
      createWritable: async () => ({ write: async (b) => { const u = new Uint8Array(await b.arrayBuffer()); window.__saved.push([o.suggestedName, await b.text(), [...u]]); }, close: async () => {} }) });""")
    block("Base", 120, 80, 10)
    block("Lid", 60, 40, 4)
    block("Bracket", 30, 20, 5)
    call("set_material", material="aluminium")
    call("new_model", name="Pin")
    call("add_feature", type="sketch", plane="XY", shapes=[{"circle": {"x": 0, "y": 0, "d": 6}}])
    call("add_feature", type="extrude", sketch="F1", d=12)
    call("new_assembly", name="Lid unit")
    for p in ("Lid", "Pin", "Pin", "Pin"):
        call("insert_part", part=p)
    print("a bought screw:", call("library_add", item="iso4762", size="M3", length=10, to="assembly")[:60])
    call("insert_part", part=SCREW)
    call("new_assembly", name="Hinge")
    call("insert_part", part="Bracket")
    call("insert_part", part="Pin")
    call("new_assembly", name="Machine")
    for p in ("Base", "Lid unit", "Lid unit", "Hinge", SCREW, SCREW, SCREW, SCREW):
        call("insert_part", part=p)
    wait(f"return Object.values({A}.bodies).every(b => b.body || b.error) && {A}.doc.instances.length === 8")

    # the AI: the summary, the tree, a sub-assembly alone
    s = json.loads(call("get_bom", mode="summary"))
    tot = {r["name"]: r["total"] for r in s["rows"]}
    print("summary:", tot, "| pieces", s["pieces"], "| mass", s["mass_g"], "g")
    check(tot == {"Base": 1, "Lid": 2, "Pin": 7, "Bracket": 1, SCREW: 8}, "Machine's totals: Base 1, Lid 2, Pin 7 (3 × 2 + 1), Bracket 1, screw 8 (4 + 2 × 2)")
    check(s["pieces"] == 19 and s["parts"] == 5, "19 pieces, 5 different parts")
    pin = next(r for r in s["rows"] if r["name"] == "Pin")
    print("where the Pins come from:", pin["from"])
    check(pin["from"] == "Lid unit: 3 × 2 + Hinge: 1", "Pin's count: Lid unit 3 × 2 + Hinge 1")
    scr = next(r for r in s["rows"] if r["name"] == SCREW)
    check(scr["make_buy"] == "bought" and scr.get("pn") and "Machine: 4" in scr["from"], f"the screw is bought ({scr.get('pn')}), 4 of them directly in Machine")
    mb = {r["name"]: r["make_buy"] for r in s["rows"]}
    check(mb["Base"] == "made (3D print)" and mb["Bracket"] == "made", f"make / buy by material (aluminium: made, not printed): {mb}")
    check(all(r.get("mass_g_each") for r in s["rows"]), "every row has its mass")
    st = json.loads(call("get_bom", mode="structured"))
    tree = [(r["level"], r["name"], r["qty_per_parent"], r["total"]) for r in st["rows"]]
    print("structured:", tree)
    check(tree == [(1, "Base", 1, 1), (1, "Lid unit", 2, 2), (2, SCREW, 2, 4), (2, "Lid", 1, 2), (2, "Pin", 3, 6), (1, "Hinge", 1, 1),
                   (2, "Pin", 1, 1), (2, "Bracket", 1, 1), (1, SCREW, 4, 4)], "the tree (by position in each assembly): qty per parent and the running total")
    sub = json.loads(call("get_bom", assembly="Lid unit"))
    check({r["name"]: r["total"] for r in sub["rows"]} == {"Lid": 1, "Pin": 3, SCREW: 2}, "Lid unit alone: Lid 1, Pin 3, screw 2")
    check(call("get_bom", assembly="Nope").startswith("ERROR"), "an unknown assembly is refused")
    pl = {r["name"]: (r["qty"], r.get("total")) for r in json.loads(call("get_assembly"))["parts_list"]}
    print("get_assembly parts_list (qty here, total):", pl)
    check(pl[SCREW] == (4, 8) and pl["Base"] == (1, 1) and pl["Lid unit"][1] is None, "parts_list `total`: the screw 4 here, 8 in all")
    pos = {r["name"]: r["pos"] for r in s["rows"]}

    # 📋 BOM in the panel
    js(f"{A}.sel = null; {A}.render()")
    time.sleep(0.3)
    d.find_element(By.CSS_SELECTOR, "#asmRight [data-a=bomWin]").click()
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 5")
    rows = table()
    print("📋 Summary:", [r[:3] for r in rows])
    check([(r[1], r[2]) for r in rows] == [("Base", "1"), ("Lid", "2"), ("Pin", "7"), (SCREW, "8"), ("Bracket", "1")] or
          sorted((r[1], r[2]) for r in rows) == sorted([("Base", "1"), ("Lid", "2"), ("Pin", "7"), (SCREW, "8"), ("Bracket", "1")]), "the Summary table's totals")
    check(all(r[0] == str(pos[r[1]]) for r in rows), "positions as the AI's (the drawing set's numbers)")
    print("   ", d.find_element(By.CSS_SELECTOR, ".bom-sum").text)
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-mode=structured]").click()
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 9")
    print("📋 Structured:", [r[1:4] for r in table()])
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-fold]").click()  # Lid unit folded
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 6")
    check(True, "a sub-assembly folds (9 -> 6 rows)")
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-fold]").click()
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 9")
    d.save_screenshot(f"{OUT}/bom_structured.png")
    # ⬇ CSV of the tree, then the summary's
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-a=csv]").click()
    wait("return window.__saved.length === 1", 20)
    name, text, head = js("return window.__saved[0]")
    lines = text.lstrip(chr(0xFEFF)).split("\r\n")
    print("CSV:", name, "|", lines[0])
    check(head[:3] == [0xEF, 0xBB, 0xBF] and len(lines) == 10 and lines[0].startswith("Pos.;Level;Name;Qty per parent;Total qty"), "the structured CSV: UTF-8 mark, a header and 9 rows")
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-mode=summary]").click()
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 5")
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-a=csv]").click()
    wait("return window.__saved.length === 2", 20)
    lines = js("return window.__saved[1][1]").lstrip(chr(0xFEFF)).split("\r\n")
    for ln in lines:
        print("   ", ln[:110])
    pin_line = next(ln for ln in lines if ";Pin;" in ln).split(";")
    check(pin_line[2] == "7" and pin_line[-1] == "made (3D print)", "the summary CSV: Pin 7, made (3D print)")
    # ⬇ Excel
    xl = d.find_element(By.CSS_SELECTOR, ".bom-win [data-a=xlsx]")
    if xl.get_attribute("disabled"):
        print("   Excel: not installed on the server (CSV only):", d.find_element(By.CSS_SELECTOR, ".bom-win .row-actions .hint").text)
    else:
        xl.click()
        wait("return window.__saved.length === 3", 30)
        name, _, raw = js("return window.__saved[2]")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(bytes(raw)))
        ws = wb[wb.sheetnames[0]]
        print("Excel:", name, wb.sheetnames, [c.value for c in ws[1]][:4])
        check(len(wb.sheetnames) == 2 and ws.max_row == 6, "the workbook: Summary + Structured, 5 rows")
    # another assembly of the hierarchy: Hinge
    sel = d.find_element(By.CSS_SELECTOR, ".bom-win .bom-of")
    print("assemblies to choose:", [o.text.strip() for o in sel.find_elements(By.TAG_NAME, "option")])
    js("const s = document.querySelector('.bom-win .bom-of'); s.value = [...s.options].find(o => o.text.includes('Hinge')).value; s.dispatchEvent(new Event('change'))")
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 2")
    check(sorted(r[1] for r in table()) == ["Bracket", "Pin"], "Hinge's own BOM chosen in the window")

    # in Czech
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-a=close]").click()
    d.find_element(By.ID, "langBtn").click()
    time.sleep(0.4)
    js(f"{A}.openBom()")
    wait("return document.querySelectorAll('.bom-win tbody tr').length === 5")
    time.sleep(0.4)
    heads = js("return [...document.querySelectorAll('.bom-win thead th')].map(t => t.innerText.trim())")
    print("Czech headers:", heads, "|", d.find_element(By.CSS_SELECTOR, ".bom-sum").text)
    check(heads[:4] == ["Poz.", "Název", "Ks", "Materiál"] and "Výroba / nákup" in heads, "the headers in Czech")
    check("nakupovaný" in [r[-1] for r in table()] and "vyráběný (3D tisk)" in [r[-1] for r in table()], "make / buy in Czech")
    mats = {r[1]: r[3] for r in table()}
    print("Czech materials:", mats)
    check(mats["Bracket"] == "hliník" and mats["Base"] == "PLA" and not any(m in EN_MATS for m in mats.values()), "materials by their Czech names (hliník), none in English")
    n0 = len(js("return window.__saved"))
    d.find_element(By.CSS_SELECTOR, ".bom-win [data-a=csv]").click()
    wait(f"return window.__saved.length === {n0 + 1}", 20)
    lines = js(f"return window.__saved[{n0}][1]").lstrip(chr(0xFEFF)).split("\r\n")
    print("Czech CSV:", lines[0], "|", next(ln for ln in lines if ";Bracket;" in ln))
    check(lines[0].startswith("Poz.;Název;Ks;Materiál;Hmotnost/ks (g)") and lines[0].endswith("Výroba / nákup")
          and ";hliník;" in next(ln for ln in lines if ";Bracket;" in ln) and not any(f";{m};" in ln for ln in lines for m in EN_MATS), "the Czech CSV: headers and materials in Czech")
    time.sleep(0.3)
    print("toasts:", js("return window.__toasts.slice(-1)"))
    check(js("return window.__toasts.at(-1)") == "Kusovník uložen: 5 řádků.", "the BOM saved toast in Czech")
    js("document.querySelector('.bom-win [data-a=close]').click()")
    before = set(js(f"return {A}.doc.instances.map(i => i.id)"))
    call("library_add", item="iso4762", size="M3", length=10, to="assembly")
    time.sleep(0.3)
    print("toasts:", js("return window.__toasts.slice(-1)"))
    check(js("return window.__toasts.at(-1)") == f"{SCREW} je v sestavě — připojte ho přes 🔗 Vazba.", "the library part's toast in Czech")
    for i in set(js(f"return {A}.doc.instances.map(i => i.id)")) - before:
        call("remove_from_assembly", part=i)
    wait(f"return {A}.doc.instances.length === 8")
    js(f"{A}.openBom()")
    wait("return document.querySelectorAll('.bom-win tbody tr').length >= 5")
    d.save_screenshot(f"{OUT}/bom_cs.png")
    js("document.querySelector('.bom-win [data-mode=structured]').click()")
    time.sleep(0.5)
    d.save_screenshot(f"{OUT}/bom_cs_structured.png")
    js("document.querySelector('.bom-win [data-a=close]').click()")
    d.find_element(By.ID, "langBtn").click()
    time.sleep(0.3)

    # the Machine's drawing sheet: 📋 Add BOM page
    mid = js("return window.__app.project.parts.find(p => p.name === 'Machine').id")
    print("sheet:", call("sheet_open", part=mid)[:80])
    wait("return !!window.__app.sheet.d?.built && !!window.__app.sheet.d?.asm", 120)
    call("sheet_edit", template="Modern")  # (this test reads the Modern parts list's texts; the default is B — ISO 7200)
    wait("return !window.__app.sheet.busy", 60)
    d.find_element(By.CSS_SELECTOR, "#shRight [data-a=bomAdd]").click()
    wait("return !!document.querySelector('.sh-sheetpage[data-bom]')", 120)
    time.sleep(0.6)
    d.save_screenshot(f"{OUT}/bom_sheet_set.png")
    check(js("return window.__app.sheet.d.bomPage") == "summary", "the sheet keeps bomPage")
    balloons = js("""const S = window.__app.sheet, P = window.__app.project, ctx = window.__app.sheet.hooks.drawingCtx();
        const id = (n) => P.parts.find(q => q.name === n).id;
        return Promise.all(['Machine', 'Lid unit', 'Hinge'].map(n => S.assemblyPage(id(n), ctx).then(r => r.rows.map(x => [x.name, x.pos]))))""")
    nums = {}
    for page in balloons:
        for nm, p in page:
            nums.setdefault(nm, set()).add(p)
    print("the balloons' numbers on the assembly pages:", {k: sorted(v) for k, v in nums.items()})
    check(all(len(v) == 1 for v in nums.values()) and all(pos[n] in nums[n] for n in pos), "the BOM's positions are the balloons' numbers on every assembly page")
    # the drawing set's PDF (the AI's export of the assembly)
    print("open Machine:", call("open_part", id=mid)[:80])
    wait(f"return {A}.active && {A}.doc.instances.length === 8")
    out = json.loads(js("return window.__app.agentHandlers.export_data({format: 'pdf'}).then(r => JSON.stringify(r.data.pages))"))
    print("pages:", out)
    check(any(p.startswith("📋") for p in out), "the drawing set has the BOM page")
    res = js("""return window.__app.agentHandlers.export_data({format: 'pdf'}).then(r => fetch('/api/sheet/pdf', {method: 'POST', headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({svgs: r.data.svgs, name: 'set'})})).then(x => x.arrayBuffer()).then(b => [...new Uint8Array(b)])""")
    path = os.path.join(OUT, "bom_set.pdf")
    with open(path, "wb") as f:
        f.write(bytes(res))
    doc = fitz.open(path)
    bom_i = next(i for i, p in enumerate(doc) if "Bill of materials" in p.get_text() and "Make / buy" in p.get_text())
    page = doc[bom_i]
    page.get_pixmap(dpi=90).save(f"{OUT}/bom_page.png")
    words = page.get_text("words")
    found = {}
    for nm in pos:
        first = nm.split()[0]
        w = [x for x in words if x[4] == first and x[1] > 25]
        if not w:
            continue
        y = w[0][1]
        left = sorted((x for x in words if abs(x[1] - y) < 1 and x[0] < w[0][0]), key=lambda x: x[0])
        found[nm] = int(left[0][4]) if left and left[0][4].isdigit() else None
    print(f"PDF page {bom_i + 1} / {len(doc)}: positions", found)
    check(found == pos, "the PDF's BOM page: the same positions as the balloons")
    # a long BOM flows on
    longer = js("""const S = window.__app.sheet, P = window.__app.project, base = S.hooks.partData, kind = S.hooks.kindOf, name = S.hooks.partName;
        const fakes = Array.from({length: 40}, (_, k) => 'FAKE' + k), id = P.parts.find(q => q.name === 'Machine').id;
        S.hooks.partData = (q) => (q === id ? {instances: [...base(q).instances, ...fakes.map((f, k) => ({id: 'Z' + k, part: f}))]} : fakes.includes(q) ? {material: 'PLA'} : base(q));
        S.hooks.kindOf = (q) => (fakes.includes(q) ? 'model' : kind(q));
        S.hooks.partName = (q) => (fakes.includes(q) ? 'Spacer ' + q.slice(4) : name(q));
        try { return S.bomPages(id, {bomPage: 'summary'}, S.d.built).map(p => [p.bom.lines.length, p.bom.page, p.bom.pages, p.bom.more]); }
        finally { S.hooks.partData = base; S.hooks.kindOf = kind; S.hooks.partName = name; }""")
    print("45 parts: pages", longer)
    check(len(longer) >= 2 and sum(p[0] for p in longer) == 45 and longer[0][3] and not longer[-1][3], "a long BOM flows onto more pages")
    # the drawing set in Czech: the BOM page and the assembly pages' parts lists
    d.find_element(By.ID, "langBtn").click()
    time.sleep(0.4)
    cz = json.loads(js("return window.__app.agentHandlers.export_data({format: 'pdf'}).then(r => JSON.stringify({pages: r.data.pages, svgs: r.data.svgs}))"))
    texts = [re.findall(r"<text[^>]*>([^<]*)</text>", s) for s in cz["svgs"]]
    bom_t = next(tx for p, tx in zip(cz["pages"], texts) if p.startswith("📋"))
    every = {x.strip() for tx in texts for x in tx}
    print("Czech BOM page:", bom_t[:12])
    check(any(x.startswith("Kusovník") for x in bom_t) and "Výroba / nákup" in bom_t and "Materiál" in bom_t and "hliník" in bom_t, "the BOM page in Czech: its title, headers and materials")
    check(not every & EN_MATS and not every & {"Make / buy", "Material", "Qty", "Part number", "made (3D print)", "bought"}, f"no English material names or headers on any page ({sorted(every & (EN_MATS | {'Material', 'Make / buy'}))})")
    d.find_element(By.ID, "langBtn").click()
finally:
    print("console errors:", [e["message"][:200] for e in d.get_log("browser") if e["level"] == "SEVERE"])
    d.quit()
print("FAILED:" if bad else "ALL OK", bad)
